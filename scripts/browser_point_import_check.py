"""Device information and simplified point import against the served production UI."""

import socket

from openpyxl import load_workbook
from playwright.async_api import expect
from pymodbus.client import ModbusTcpClient

from simulator.models import Device, Point
from simulator.point_excel import COLUMNS


async def check_point_import(page, url, output):
    with socket.socket() as available:
        available.bind(("127.0.0.1", 0))
        port = available.getsockname()[1]
    d = Device(
        id="simple-import-device",
        name="点位模板验收",
        port=port,
        description="设备说明\n第二行",
        read_identity=True,
        points=[
            Point(
                id="simple-existing",
                name="模板更新",
                type="Float32",
                initial=12,
                writable=True,
            ),
            Point(id="simple-keep", name="未导入保留", address=8, initial=99),
        ],
    )
    config = await (await page.request.get(url + "/api/config")).json()
    config["devices"].append(d.model_dump())
    response = await page.request.put(url + "/api/config", data=config)
    assert response.status == 200, await response.text()
    await page.set_viewport_size({"width": 1440, "height": 1000})
    await page.goto(url)
    card = page.locator(".device-link").filter(has_text=d.name)
    await card.click()
    await card.click(button="right")
    menu = page.get_by_role("menu", name="设备操作", exact=True)
    await menu.get_by_role("menuitem", name="设备信息", exact=True).click()
    panel = page.get_by_role("region", name="设备信息", exact=True)
    await expect(panel).to_be_visible()
    fields = await panel.locator("dt").all_text_contents()
    assert fields[0] == "设备 ID" and len(fields) == 22, fields
    await expect(panel.locator("dd").first.locator("code")).to_have_text(d.id)
    await expect(panel).to_contain_text("第二行")
    await expect(panel).to_contain_text("写多个保持寄存器")
    await expect(panel).to_contain_text(d.identity.vendor)
    for width in (1440, 1024, 390):
        await page.set_viewport_size({"width": width, "height": 844})
        assert await page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        await page.screenshot(
            path=str(output / f"device-information-{width}.png"), full_page=True
        )
    await page.set_viewport_size({"width": 1440, "height": 1000})
    await page.get_by_role("button", name="点位监控", exact=True).click()
    response = await page.request.post(url + f"/api/devices/{d.id}/actions/start")
    assert response.status == 200
    protocol = ModbusTcpClient("127.0.0.1", port=port, timeout=2)
    try:
        assert protocol.connect()
        assert not protocol.write_registers(
            0, [0x4216, 0], device_id=1
        ).isError()  # Float32 37.5, hold.
        await page.get_by_role("button", name="导入", exact=True).click()
        drawer = page.get_by_role("dialog", name="Excel 导入", exact=True)
        async with page.expect_download() as download_info:
            await drawer.get_by_role("button", name="下载模板", exact=True).click()
        path = output / "single-sheet-template.xlsx"
        await (await download_info.value).save_as(str(path))
        wb = load_workbook(path)
        assert wb.sheetnames == ["点位"]
        ws = wb["点位"]
        assert [c.value for c in ws[1]] == list(COLUMNS)
        rows = [
            {
                "设备 ID": d.id,
                "名称": "模板更新",
                "分组": "导入组",
                "数据区": "保持寄存器",
                "协议地址": 0,
                "类型": "Float32",
                "初始值": 20,
                "主机可写": "是",
                "最小值": 10,
                "最大值": 30,
            },
            {
                "设备 ID": d.id,
                "名称": "随机开关",
                "数据区": "线圈",
                "协议地址": 0,
                "类型": "Bool",
            },
        ]
        for row in rows:
            ws.append([row.get(c) for c in COLUMNS])
        wb.save(path)
        wb.close()
        feedback = drawer.get_by_role("status", name="操作提示", exact=True)
        await feedback.get_by_role("button", name="关闭操作提示", exact=True).click()
        await drawer.locator("input[type=file]").set_input_files(path)
        await drawer.get_by_role("button", name="解析并校验", exact=True).click()
        apply = drawer.get_by_role("button", name="应用配置", exact=True)
        await expect(apply).to_be_enabled(timeout=15000)
        await page.screenshot(path=str(output / "point-template-preview.png"))
        await apply.click()
        await expect(
            drawer.get_by_role("heading", name="导入已完成", exact=True)
        ).to_be_visible()
        await page.keyboard.press("Escape")
        points = (
            await (await page.request.get(url + f"/api/devices/{d.id}/points")).json()
        )["items"]
        a, b, c = points
        assert (
            a["id"] == "simple-existing" and a["initial"] == 20 and a["value"] == 37.5
        )
        assert b["id"] == "simple-keep" and b["value"] == 99
        assert c["name"] == "随机开关" and c["strategy"]["kind"] == "random"
        assert c["initial"] in (0, 1) and c["precision"] == 0
        assert protocol.read_holding_registers(0, count=2, device_id=1).registers == [
            0x4216,
            0,
        ]
        assert protocol.read_coils(0, count=1, device_id=1).bits[0] in (False, True)
        await page.get_by_role(
            "button", name="随机开关：查看／编辑", exact=True
        ).click()
        editor = page.get_by_role("dialog")
        await expect(editor.get_by_label("策略类型", exact=True)).to_have_value(
            "random"
        )
        await expect(editor.get_by_label("最大值", exact=True)).to_have_value("1")
        await page.keyboard.press("Escape")
    finally:
        protocol.close()
        await page.request.post(url + f"/api/devices/{d.id}/actions/stop")
    return [
        "device information renders all 22 information fields, selectable ID first, Chinese meanings at 1440/1024/390 without horizontal overflow",
        "real UI downloads one-sheet 14-column template with dropdowns, live name upsert retains stable ID/current held value/unlisted point and existing TCP connection",
        "Boolean point default random strategy imports as 0/1 and reopens correctly in editor",
    ]
