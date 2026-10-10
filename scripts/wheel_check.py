"""Verify an installed wheel outside the checkout, including real TCP I/O."""

import argparse
import re
import socket
import sys
import tempfile
from io import BytesIO
from pathlib import Path


def run(target):
    sys.path.insert(0, str(Path(target).resolve()))
    import simulator
    from openpyxl import load_workbook
    from fastapi.testclient import TestClient
    from pymodbus.client import ModbusTcpClient
    from simulator.api import create_app

    assert Path(simulator.__file__).is_relative_to(Path(target).resolve())
    with socket.socket() as available:
        available.bind(("127.0.0.1", 0))
        port = available.getsockname()[1]
    with tempfile.TemporaryDirectory() as data:
        with TestClient(create_app(Path(data))) as http:
            html = http.get("/devices/overview")
            assert html.status_code == 200 and '<div id="app">' in html.text
            assert http.get("/help").text == html.text
            assets = re.findall(r'(?:src|href)="(/assets/[^\"]+)"', html.text)
            assert assets and all(
                http.get(asset).status_code == 200 for asset in assets
            )
            assert http.get("/FONT-LICENSE.txt").status_code == 200
            fonts = list(Path(simulator.__file__).parent.glob("static/assets/*.woff2"))
            assert fonts and all(
                http.get("/assets/" + font.name).status_code == 200 for font in fonts
            )
            assert http.get("/api/health/ready").status_code == 200
            result = http.post(
                "/api/templates/thermal", json={"version": 0, "port": port}
            )
            assert result.status_code == 200, result.text
            key = result.json()["id"]
            assert http.post(f"/api/devices/{key}/actions/start").status_code == 200
            client = ModbusTcpClient("127.0.0.1", port=port, timeout=2)
            try:
                assert client.connect()
                words = [0x4296, 0x8000]  # Float32 75.25, big byte/word order.
                assert not client.write_registers(0, words, device_id=1).isError()
                assert (
                    client.read_holding_registers(0, count=2, device_id=1).registers
                    == words
                )
                rows = http.get(f"/api/devices/{key}/points").json()["items"]
                assert rows[1]["value"] == 75.25 and rows[1]["raw"] == words
                assert rows[1]["type"] == "Float32"
                assert rows[1]["scale"] == 1 and rows[1]["precision"] == 2
                assert client.write_register(65535, 1, device_id=1).isError()
                history = http.get(f"/api/devices/{key}/assignments").json()["items"]
                assert [r["outcome"] for r in history] == ["failed", "success"]
                assert all(
                    r["origin"] == "modbus" and r["host"] == "127.0.0.1"
                    for r in history
                )
                assert history[1]["changes"][0]["after"] == 75.25
            finally:
                client.close()
            # Freeze generation while checking that a compatible import retains the value.
            assert http.post(f"/api/devices/{key}/actions/pause").status_code == 200
            table = http.get(f"/api/export/point-table?device={key}")
            assert table.status_code == 200
            book = load_workbook(BytesIO(table.content))
            assert book.sheetnames == ["Modbus点表"]
            assert book.active.max_row == len(rows) + 1
            assert book.active.max_column == 20
            assert book.active.freeze_panes is None
            assert not {"设备 ID", "分组", "初始值"}.intersection(
                cell.value for cell in book.active[1]
            )
            book.close()
            template = http.get("/api/export/template")
            assert template.status_code == 200
            wb = load_workbook(BytesIO(template.content))
            assert wb.sheetnames == ["点位"] and wb["点位"].max_column == 14
            assert len(wb["点位"].data_validations.dataValidation) == 4
            values = [
                key,
                rows[1]["name"],
                "",
                "保持寄存器",
                0,
                "Float32",
                1,
                0,
                "℃",
                20,
                "是",
                "均匀随机",
                100,
                0,
            ]
            for col, value in enumerate(values, 1):
                wb["点位"].cell(2, col, value)
            stream = BytesIO()
            wb.save(stream)
            wb.close()
            preview = http.post(
                "/api/import/preview",
                files={"file": ("points.xlsx", stream.getvalue())},
            ).json()
            assert (
                not preview["errors"]
                and preview["changed"] == 1
                and preview["added"] == 0
            )
            assert (
                http.post("/api/import/apply", json=preview["config"]).status_code
                == 200
            )
            points = http.get(f"/api/devices/{key}/points").json()["items"]
            assert points[1]["id"] == rows[1]["id"] and points[1]["value"] == 75.25
            assert http.post(f"/api/devices/{key}/actions/stop").status_code == 200
        with TestClient(create_app(Path(data))) as restarted:
            assert (
                restarted.get(f"/api/devices/{key}/assignments").json()["items"]
                == history
            )
            target_id = rows[1]["id"]
            assert (
                restarted.post(
                    f"/api/devices/{key}/assign",
                    json={"origin": "web", "items": [{"id": target_id, "value": 60.5}]},
                ).status_code
                == 200
            )
            after = restarted.get(f"/api/devices/{key}/assignments").json()["items"]
            assert len(after) == 3 and after[0]["origin"] == "web"
            assert after[0]["changes"][0]["before"] == 75.25
            assert after[0]["changes"][0]["after"] == 60.5
            assert (
                restarted.post(
                    f"/api/devices/{key}/assign",
                    json={"items": [{"id": target_id, "value": "invalid"}]},
                ).status_code
                == 422
            )
            assert (
                restarted.post(f"/api/devices/{key}/actions/start").status_code == 200
            )
            client = ModbusTcpClient("127.0.0.1", port=port, timeout=2)
            try:
                assert client.connect()
                assert not client.write_registers(
                    0, [0x4298, 0x8000], device_id=1
                ).isError()
                assert client.write_register(65535, 1, device_id=1).isError()
            finally:
                client.close()
            after = restarted.get(f"/api/devices/{key}/assignments").json()["items"]
            assert len(after) == 6
            assert [row["outcome"] for row in after[:4]] == [
                "failed",
                "success",
                "failed",
                "success",
            ]
            assert after[1]["changes"][0]["after"] == 76.25
        with TestClient(create_app(Path(data))) as again:
            assert again.get(f"/api/devices/{key}/assignments").json()["items"] == after
    print(
        "Installed wheel passed: SPA/help, local assets/fonts, API, real Modbus write/read, single-sheet template/name upsert, new success/failure assignments after snapshot restore and persistence"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--target", required=True)
    run(parser.parse_args().target)
