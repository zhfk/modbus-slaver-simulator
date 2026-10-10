"""Run the real frozen binary outside the checkout with three TCP devices."""

import argparse
import contextlib
import asyncio
from datetime import datetime, timezone
import json
from io import BytesIO
import os
from pathlib import Path
import re
import signal
import socket
import struct
import subprocess
import tempfile
import time

import httpx
import psutil
from openpyxl import load_workbook
from pymodbus.client import ModbusTcpClient
from websockets.asyncio.client import connect


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def stop_process(process):
    children = (
        psutil.Process(process.pid).children(recursive=True)
        if process.poll() is None
        else []
    )
    if process.poll() is None:
        if os.name == "nt":
            process.send_signal(signal.CTRL_BREAK_EVENT)
        else:
            process.terminate()
        try:
            process.wait(timeout=25)
        except subprocess.TimeoutExpired:
            for child in children:
                with contextlib.suppress(psutil.NoSuchProcess):
                    child.kill()
            process.kill()
            process.wait(timeout=5)
    for child in children:
        if child.is_running() and child.status() != psutil.STATUS_ZOMBIE:
            child.wait(timeout=5)


def check(binary, output):
    binary = Path(binary).resolve()
    version = subprocess.check_output(
        [str(binary), "--version"], text=True, timeout=10
    ).strip()
    checks = []
    with tempfile.TemporaryDirectory(prefix="modbus-release-check-") as temporary:
        cwd = Path(temporary)
        data, http_port, tcp_port = cwd / "data # 中文", free_port(), free_port()
        env = os.environ.copy()
        for key in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV"):
            env.pop(key, None)
        # Application and Excel children must rely on the bundled interpreter.
        env["PATH"] = str(cwd)
        with (cwd / "service.log").open("wb") as log:
            process = subprocess.Popen(
                [
                    str(binary),
                    "--data-dir",
                    str(data),
                    "--port",
                    str(http_port),
                    "--interval",
                    "0.5",
                ],
                cwd=cwd,
                env=env,
                stdout=log,
                stderr=subprocess.STDOUT,
                creationflags=subprocess.CREATE_NEW_PROCESS_GROUP
                if os.name == "nt"
                else 0,
            )
        address = f"http://127.0.0.1:{http_port}"
        try:
            with httpx.Client(base_url=address, trust_env=False, timeout=10) as http:
                deadline = time.monotonic() + 40
                while True:
                    if process.poll() is not None:
                        raise RuntimeError(
                            (cwd / "service.log").read_text(errors="replace")
                        )
                    try:
                        if http.get("/api/health/ready").status_code == 200:
                            break
                    except httpx.HTTPError:
                        pass
                    if time.monotonic() > deadline:
                        raise TimeoutError("Frozen application did not become ready")
                    time.sleep(0.1)
                assert http.get("/api/health").json()["storage"]["data_dir"] == str(
                    data.resolve()
                )
                html = http.get("/devices/overview")
                assert html.status_code == 200 and '<div id="app">' in html.text
                assets = re.findall(r'(?:src|href)="(/assets/[^\"]+)"', html.text)
                assert assets and all(
                    http.get(path).status_code == 200 for path in assets
                )
                assert http.get("/FONT-LICENSE.txt").status_code == 200
                assert http.get("/api/missing").status_code == 404
                checks.append(
                    "bundled page/assets/fonts, deep links and actual data directory without system Python/PATH"
                )
                keys = []
                for unit in (1, 2, 3):
                    config = http.get("/api/config").json()
                    response = http.post(
                        "/api/templates/thermal",
                        json={
                            "version": config["version"],
                            "name": f"从机 {unit}#",
                            "port": tcp_port,
                            "unit_id": unit,
                        },
                    )
                    response.raise_for_status()
                    key = response.json()["id"]
                    keys.append(key)
                    http.post(f"/api/devices/{key}/actions/start").raise_for_status()
                client = ModbusTcpClient("127.0.0.1", port=tcp_port, timeout=2)
                try:
                    assert client.connect()
                    for unit, key in zip((1, 2, 3), keys):
                        value = 60.25 + unit
                        words = list(struct.unpack(">HH", struct.pack(">f", value)))
                        assert not client.write_registers(
                            0, words, device_id=unit
                        ).isError()
                        assert (
                            client.read_holding_registers(
                                0, count=2, device_id=unit
                            ).registers
                            == words
                        )
                        rows = http.get(f"/api/devices/{key}/points").json()["items"]
                        assert rows[1]["raw"] == words
                        assert rows[1]["value"] == value
                        assert rows[1]["type"] == "Float32"
                        assert rows[1]["scale"] == 1 and rows[1]["precision"] == 2
                    http.post(f"/api/devices/{keys[0]}/actions/stop").raise_for_status()
                    for unit in (2, 3):
                        assert client.read_holding_registers(
                            0, count=2, device_id=unit
                        ).registers == list(
                            struct.unpack(">HH", struct.pack(">f", 60.25 + unit))
                        )
                finally:
                    client.close()
                checks.append(
                    "three devices share one endpoint with isolated Unit IDs, values and stop"
                )

                async def websocket_snapshot():
                    async with connect(
                        address.replace("http://", "ws://") + "/api/live", proxy=None
                    ) as ws:
                        point = http.get(f"/api/devices/{keys[1]}/points").json()[
                            "items"
                        ][1]
                        await ws.send(
                            json.dumps({"device": keys[1], "ids": [point["id"]]})
                        )
                        snapshot = json.loads(await asyncio.wait_for(ws.recv(), 5))
                        assert (
                            snapshot["complete"]
                            and snapshot["items"][0]["raw"] == point["raw"]
                        )

                asyncio.run(websocket_snapshot())
                checks.append("real WebSocket snapshot from bundled protocol adapter")
                export = http.get("/api/export/config", timeout=65)
                export.raise_for_status()
                assert export.content.startswith(b"PK")
                preview = http.post(
                    "/api/import/preview",
                    files={"file": ("config.xlsx", export.content)},
                    timeout=65,
                )
                preview.raise_for_status()
                assert not preview.json()["errors"] and preview.json()["changed"] == 12
                assert not list((data / "tmp").iterdir())
                checks.append(
                    "spawned Excel export/import workers and temporary cleanup"
                )
                table = http.get("/api/export/point-table", timeout=65)
                table.raise_for_status()
                book = load_workbook(BytesIO(table.content))
                assert book.sheetnames == ["Modbus点表"]
                assert book.active.max_row == 13 and book.active.max_column == 20
                headers = [cell.value for cell in book.active[1]]
                assert not {"设备 ID", "分组", "初始值"}.intersection(headers)
                assert (
                    book.active.freeze_panes is None
                    and book.active.sheet_view.pane is None
                )
                assert {
                    row[3] for row in book.active.iter_rows(min_row=2, values_only=True)
                } == {1, 2, 3}
                book.close()
                template = http.get("/api/export/template", timeout=65)
                template.raise_for_status()
                book = load_workbook(BytesIO(template.content))
                assert book.sheetnames == ["点位"] and book.active.max_column == 14
                assert book.active.sheet_view.selection[0].activeCell == "A2"
                assert len(book.active.data_validations.dataValidation) == 4
                book.close()
                checks.append(
                    "single-pane 20-column host point table and 14-column import template from frozen Excel workers"
                )
                http.post(f"/api/devices/{keys[1]}/actions/pause").raise_for_status()
                response = http.get("/api/health/live")
                response.raise_for_status()
                assert response.json()["live"]
                checks.append("external supervisor observes real business progress")
        except BaseException:
            print((cwd / "service.log").read_text(errors="replace"))
            raise
        finally:
            stop_process(process)
        with (cwd / "restart.log").open("wb") as log:
            restarted = subprocess.Popen(
                [
                    str(binary),
                    "serve",
                    "--data-dir",
                    str(data),
                    "--port",
                    str(http_port),
                ],
                cwd=cwd,
                env=env,
                stdout=log,
                stderr=subprocess.STDOUT,
                creationflags=subprocess.CREATE_NEW_PROCESS_GROUP
                if os.name == "nt"
                else 0,
            )
        try:
            with httpx.Client(base_url=address, trust_env=False, timeout=10) as http:
                deadline = time.monotonic() + 40
                while True:
                    if restarted.poll() is not None:
                        raise RuntimeError(
                            (cwd / "restart.log").read_text(errors="replace")
                        )
                    try:
                        if http.get("/api/health/ready").status_code == 200:
                            break
                    except httpx.HTTPError:
                        pass
                    if time.monotonic() > deadline:
                        raise TimeoutError("Frozen restart readiness timeout")
                    time.sleep(0.1)
                states = {d["id"]: d for d in http.get("/api/devices").json()}
                assert states[keys[0]]["status"] == "stopped"
                assert (
                    states[keys[1]]["status"] == "running" and states[keys[1]]["paused"]
                )
                assert (
                    states[keys[2]]["status"] == "running"
                    and not states[keys[2]]["paused"]
                )
                client = ModbusTcpClient("127.0.0.1", port=tcp_port, timeout=2)
                try:
                    assert client.connect()
                    assert client.read_holding_registers(
                        0, count=2, device_id=2
                    ).registers == list(struct.unpack(">HH", struct.pack(">f", 62.25)))
                finally:
                    client.close()
                checks.append(
                    "same-directory frozen process restart restores running/stopped/paused state and actual Float32 TCP values"
                )
        except BaseException:
            print((cwd / "restart.log").read_text(errors="replace"))
            raise
        finally:
            stop_process(restarted)
        # The same binary provides maintenance without an installed Python.
        for command in (
            ["compact"],
            ["restore", "--backup", str(data / "backups/config.0.db")],
        ):
            subprocess.run(
                [str(binary), "maintenance", "--data-dir", str(data), *command],
                cwd=cwd,
                env=env,
                check=True,
                timeout=20,
            )
        checks.append("offline compact and backup restore with bundled SQLite")
    report = {
        "passed": True,
        "version": version,
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "checks": checks,
    }
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--binary", required=True)
    parser.add_argument("--output", default="artifacts/release-check.json")
    args = parser.parse_args()
    check(args.binary, args.output)
