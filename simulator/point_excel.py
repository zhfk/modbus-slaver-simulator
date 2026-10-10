"""Single-sheet point templates, parsed inside the bounded Excel worker."""

import json
import math
import random

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.views import Selection
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .models import Configuration, Point, uid

COLUMNS = {
    "设备 ID": "device_id",
    "名称": "name",
    "分组": "group",
    "数据区": "area",
    "协议地址": "address",
    "类型": "type",
    "倍率": "scale",
    "偏移": "offset",
    "单位": "unit",
    "初始值": "initial",
    "主机可写": "writable",
    "策略类型": "kind",
    "最大值": "max",
    "最小值": "min",
}
AREAS = {
    "线圈": "coil",
    "离散输入": "discrete",
    "保持寄存器": "holding",
    "输入寄存器": "input",
}
STRATEGIES = {
    "均匀随机": "random",
    "无策略": "none",
    "固定值": "fixed",
    "正弦波": "sine",
    "斜坡": "ramp",
    "随机游走": "walk",
}
LIMITS = {
    "Int16": (-32768, 32767),
    "UInt16": (0, 65535),
    "Int32": (-2147483648, 2147483647),
    "UInt32": (0, 4294967295),
    "Float32": (-3.4028234663852886e38, 3.4028234663852886e38),
    "Float64": (-1.7976931348623157e308, 1.7976931348623157e308),
}


def freeze_header(ws):
    ws.freeze_panes = "A2"
    # openpyxl otherwise leaves A1 selected in the bottom pane. Excel/WPS
    # can scroll it back to row 1 and display the frozen header twice.
    ws.sheet_view.topLeftCell = "A1"
    ws.sheet_view.selection = [
        Selection(pane="bottomLeft", activeCell="A2", sqref="A2")
    ]


def write_template(wb):
    ws = wb.active
    ws.title = "点位"
    ws.append(list(COLUMNS))
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 24
    notes = {
        "设备 ID": "必填；从设备卡片菜单的设备信息顶部复制已有设备 ID，每行填写对应设备 ID。",
        "名称": "必填；同一设备按名称精确匹配，存在则更新并保留 ID，不存在则自动生成 ID 新增。",
        "分组": "可空，默认无分组。",
        "数据区": "必填，下拉选择；离散输入和输入寄存器只能供主机读取。",
        "协议地址": "必填，从 0 开始；Float32/Int32/UInt32 占 2 个寄存器，Float64 占 4 个。",
        "类型": "必填，下拉选择；线圈/离散输入仅支持 Bool，其余区使用数值类型。",
        "倍率": "可空，默认 1，不能为 0。工程值 = 原始值 × 倍率 + 偏移；Bool 仅允许 1。",
        "偏移": "可空，默认 0；Bool 仅允许 0。",
        "单位": "可空，默认无单位。",
        "初始值": "可空；预览时按类型与最小/最大值生成可编码的随机初值，应用使用同一预览值。",
        "主机可写": "可空，默认否，下拉选择是/否。输入区不得选择是。写入模式默认保持。",
        "策略类型": "可空，默认均匀随机；默认启用、每 1 秒更新、种子 1。Bool 支持无策略、固定值、均匀随机；复杂联动在页面编辑。",
        "最大值": "可空，默认 100，按类型/倍率/偏移限制在可表示范围；Bool 默认 1。",
        "最小值": "可空，默认 0，按类型/倍率/偏移限制在可表示范围；Bool 默认 0。不得大于最大值。",
    }
    for cell in ws[1]:
        cell.comment = Comment(notes[cell.value], "Modbus Simulator")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245C86")
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        ws.column_dimensions[cell.column_letter].width = (
            24 if cell.value in ("设备 ID", "名称") else 16
        )
        # Give the first editable row explicit formats. Column defaults cover
        # later rows, while Excel/WPS also see a real input range on opening.
        number_format = (
            "@"
            if COLUMNS[cell.value]
            in (
                "device_id",
                "name",
                "group",
                "area",
                "type",
                "unit",
                "writable",
                "kind",
            )
            else "General"
        )
        ws.column_dimensions[cell.column_letter].number_format = number_format
        entry = ws.cell(2, cell.column)
        entry.number_format = number_format
        entry.alignment = Alignment(vertical="center")
        entry.fill = PatternFill("solid", fgColor="F5F7FA")
    for header, choices in {
        "数据区": list(AREAS),
        "类型": ["Bool", *LIMITS],
        "主机可写": ["是", "否"],
        "策略类型": list(STRATEGIES),
    }.items():
        col = ws.cell(1, list(COLUMNS).index(header) + 1).column_letter
        rule = DataValidation(
            type="list",
            formula1='"' + ",".join(choices) + '"',
            allow_blank=True,
            showDropDown=False,
            showInputMessage=True,
        )
        rule.promptTitle = header
        rule.prompt = "从第 2 行开始填写，点击单元格右侧箭头选择：" + "、".join(choices)
        rule.errorTitle, rule.error = (
            "请选择下拉选项",
            "请输入列表中的值，后台会再次校验。",
        )
        rule.showErrorMessage, rule.errorStyle = True, "stop"
        ws.add_data_validation(rule)
        rule.add(f"{col}2:{col}10001")
    freeze_header(ws)


class CellError(ValueError):
    def __init__(self, field, message):
        super().__init__(message)
        self.field = field


def number(item, key, default):
    field = next(h for h, k in COLUMNS.items() if k == key)
    try:
        value = float(item.get(key, default))
    except (ValueError, TypeError, OverflowError) as exc:
        raise CellError(field, f"{field} 必须为有限数值") from exc
    if not math.isfinite(value):
        raise CellError(field, f"{field} 必须为有限数值")
    return value


def build_point(item):
    """Normalize the narrow template; omitted advanced fields use model defaults."""
    for required in ("area", "address", "type"):
        if required not in item:
            raise CellError(
                next(h for h, k in COLUMNS.items() if k == required), "必填项不能为空"
            )
    if item["area"] not in AREAS:
        raise CellError("数据区", "数据区请选择：" + "、".join(AREAS))
    item["area"] = AREAS[item["area"]]
    kind = item["type"]
    if kind not in ["Bool", *LIMITS]:
        raise CellError("类型", "类型请选择模板提供的数据类型")
    writable = item.get("writable", "否")
    if writable not in ("是", "否", True, False, 0, 1):
        raise CellError("主机可写", "主机可写请选择是／否")
    item["writable"] = writable in ("是", True, 1)
    strategy = item.pop("kind", "均匀随机")
    if strategy not in STRATEGIES:
        raise CellError(
            "策略类型", "策略类型请选择模板提供的策略，复杂策略请在页面编辑"
        )
    strategy = STRATEGIES[strategy]
    if kind == "Bool" and strategy not in ("none", "fixed", "random"):
        raise CellError("策略类型", "Bool 仅支持无策略、固定值或均匀随机")
    for attr, default in (("scale", 1), ("offset", 0)):
        item[attr] = number(item, attr, default)
        if not math.isfinite(item[attr]) or (attr == "scale" and item[attr] == 0):
            raise CellError("倍率", "倍率不能为零")
    scale, offset = item["scale"], item["offset"]
    if kind == "Bool":
        low, high = 0, 1
    else:
        a, b = LIMITS[kind]
        low, high = sorted((a * scale + offset, b * scale + offset))
    lo = number(item, "min", max(low, min(0, high)))
    item.pop("min", None)
    hi = number(item, "max", min(high, max(1 if kind == "Bool" else 100, low)))
    item.pop("max", None)
    if not all(math.isfinite(v) for v in (lo, hi)) or not low <= lo <= hi <= high:
        raise CellError(
            "最小值／最大值", "最小值／最大值须有限、顺序正确且可由类型、倍率和偏移表示"
        )
    if kind == "Bool" and (lo not in (0, 1) or hi not in (0, 1)):
        raise CellError("最小值／最大值", "Bool 的最小值和最大值只能为 0／1")
    if "initial" not in item:
        rng = random.SystemRandom()
        if kind == "Bool":
            item["initial"] = rng.randint(int(lo), int(hi))
        elif kind.startswith("Float"):
            fraction = rng.random()
            item["initial"] = lo * fraction + hi * (1 - fraction)
        else:
            start, end = sorted(((lo - offset) / scale, (hi - offset) / scale))
            start, end = math.ceil(start), math.floor(end)
            if start > end:
                raise CellError("初始值", "指定范围中没有可表示的整数初始值")
            item["initial"] = rng.randint(start, end) * scale + offset
    params = {"min": lo, "max": hi}
    if strategy == "fixed":
        params["value"] = item["initial"]
    elif strategy == "sine":
        params.update(mean=lo / 2 + hi / 2, amplitude=hi / 2 - lo / 2, period=60)
    elif strategy == "ramp":
        params.update(start=lo, end=hi, duration=60, loop=True)
    elif strategy == "walk":
        params["step"] = hi / 100 - lo / 100
    item["strategy"] = {"kind": strategy, "params": params}
    return Point.model_validate(item).model_dump()


def parse_points(wb, current):
    errors, locations = [], {}
    ws = wb["点位"]
    if ws.max_column != len(COLUMNS) or ws.max_row > 10001:
        raise ValueError("点位须使用 14 列模板，最多 10000 个点位")
    iterator = ws.iter_rows(values_only=True)
    headers = next(iterator)
    if len(set(headers)) != len(headers) or set(headers) != set(COLUMNS):
        raise ValueError("点位列名非法、重复或缺少列，请使用最新模板")
    candidate = json.loads(json.dumps(current or {"devices": []}))
    devices = {d["id"]: d for d in candidate["devices"]}
    names = {}
    for d in candidate["devices"]:
        for i, p in enumerate(d["points"]):
            names.setdefault((d["id"], p["name"]), []).append(i)
    seen, added, changed = set(), 0, 0
    reverse = {v: k for k, v in COLUMNS.items()}
    for row, values in enumerate(iterator, 2):
        if row > 10001:
            raise ValueError("点位超过 10000 行限制")
        if all(v is None for v in values):
            continue
        item = {
            COLUMNS[h]: v for h, v in zip(headers, values) if v is not None and v != ""
        }
        field = "设备 ID"
        try:
            device_id = item.pop("device_id", None)
            if not isinstance(device_id, str) or device_id not in devices:
                raise ValueError("请填写设备信息中已有的设备 ID；点位导入不会创建设备")
            field = "名称"
            name = item.get("name")
            if not isinstance(name, str) or not name.strip():
                raise ValueError("名称必填，必须为非空文本")
            key = (device_id, name)
            if key in seen:
                raise ValueError("文件中设备 ID 和名称重复")
            seen.add(key)
            matches = names.get(key, [])
            if len(matches) > 1:
                raise ValueError("设备中存在多个同名点位，请先在页面改名消除歧义")
            device = devices[device_id]
            item["id"] = device["points"][matches[0]]["id"] if matches else uid()
            field = "点位属性"
            point = build_point(item)
            if matches:
                index = matches[0]
                device["points"][index] = point
                changed += 1
            else:
                index = len(device["points"])
                device["points"].append(point)
                added += 1
            locations[(device_id, index)] = row
        except (ValueError, TypeError, OverflowError) as exc:
            if isinstance(exc, ValidationError):
                for e in exc.errors(include_input=False, include_url=False):
                    column = (
                        reverse.get(e["loc"][0], str(e["loc"][0]))
                        if e["loc"]
                        else field
                    )
                    errors.append(
                        {
                            "sheet": "点位",
                            "row": row,
                            "field": column,
                            "message": e["msg"],
                        }
                    )
            else:
                errors.append(
                    {
                        "sheet": "点位",
                        "row": row,
                        "field": getattr(exc, "field", field),
                        "message": str(exc)[:512],
                    }
                )
    if not errors:
        if not seen:
            errors.append(
                {
                    "sheet": "点位",
                    "row": 2,
                    "field": "名称",
                    "message": "文件中没有点位，请填写模板后导入",
                }
            )
        else:
            try:
                candidate = Configuration.model_validate(candidate).model_dump()
            except ValidationError as exc:
                for e in exc.errors(include_input=False, include_url=False):
                    loc = e["loc"]
                    device_id = (
                        candidate["devices"][loc[1]]["id"]
                        if len(loc) > 1 and loc[0] == "devices"
                        else None
                    )
                    # Device-wide overlap/dependency checks point to the first edited row.
                    row = (
                        locations.get((device_id, loc[3]), 0)
                        if len(loc) > 3
                        else next(
                            (r for (d, _), r in locations.items() if d == device_id), 0
                        )
                    )
                    errors.append(
                        {
                            "sheet": "点位",
                            "row": row,
                            "field": "配置",
                            "message": e["msg"],
                        }
                    )
    return (
        {"errors": errors}
        if errors
        else {
            "errors": [],
            "config": candidate,
            "format": "points",
            "added": added,
            "changed": changed,
            "deleted": 0,
        }
    )
