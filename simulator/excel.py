"""Versioned Excel interchange, executed only in killable worker processes."""

import asyncio
import json
import multiprocessing
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from .models import Configuration, uid
from .errors import DomainError
from .files import directory_bytes
from .point_excel import freeze_header, parse_points, write_template
from .point_table import write_point_table

DEVICE_COLUMNS = {
    "设备 ID": "id",
    "名称": "name",
    "说明": "description",
    "绑定 IP": "host",
    "端口": "port",
    "Unit ID": "unit_id",
    "功能码": "functions",
    "未配置地址": "missing_address",
    "未知 Unit ID": "unknown_unit",
    "最大连接数": "max_connections",
    "空闲超时秒": "idle_seconds",
    "组帧超时秒": "frame_seconds",
    "自动启动": "auto_start",
    "故障参数": "faults",
    "有效地址范围": "valid_ranges",
    "设备身份": "identity",
    "支持读取身份": "read_identity",
}
POINT_COLUMNS = {
    "设备 ID": "device_id",
    "点位 ID": "id",
    "名称": "name",
    "分组": "group",
    "说明": "description",
    "数据区": "area",
    "协议地址": "address",
    "类型": "type",
    "字节顺序": "byte_order",
    "寄存器顺序": "word_order",
    "倍率": "scale",
    "偏移": "offset",
    "单位": "unit",
    "显示精度": "precision",
    "初始值": "initial",
    "主机可写": "writable",
    "写入模式": "write_mode",
    "覆盖秒数": "override_seconds",
    "重启恢复": "restore",
    "策略类型": "strategy_kind",
    "策略启用": "strategy_enabled",
    "更新周期秒": "strategy_interval",
    "随机种子": "strategy_seed",
    "依赖点位 ID": "strategy_dependencies",
    "固定值": "value",
    "最小值": "min",
    "最大值": "max",
    "波形周期秒": "period",
    "幅度": "amplitude",
    "均值": "mean",
    "复杂参数": "strategy_params",
}
PARAM_KEYS = ("value", "min", "max", "period", "amplitude", "mean")


def text_cells(ws):
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"


def write_workbook(payload, destination):
    wb = Workbook()
    meta = wb.active
    if payload.get("kind") == "template":
        write_template(wb)
        wb.save(destination)
        return {"path": str(destination)}
    if payload.get("kind") == "point-table":
        write_point_table(wb, payload["devices"])
        wb.save(destination)
        return {"path": str(destination)}
    meta.title = "格式"
    meta.append(["格式版本", 1])
    meta.append(["类型", payload.get("kind", "config")])
    meta.append(["地址格式", "protocol_zero"])
    meta.append(["设置", json.dumps(payload.get("settings", {}), ensure_ascii=False)])
    if payload.get("kind", "config") == "config":
        ds, ps = wb.create_sheet("设备"), wb.create_sheet("点位")
        ds.append(list(DEVICE_COLUMNS))
        ps.append(list(POINT_COLUMNS))
        for device in payload["devices"]:
            row = []
            for key in DEVICE_COLUMNS.values():
                value = device.get(key)
                row.append(
                    json.dumps(value, ensure_ascii=False)
                    if isinstance(value, (dict, list))
                    else value
                )
            ds.append(row)
            for point in device["points"]:
                flat = {**point, "device_id": device["id"]}
                strategy = point["strategy"]
                flat.update(
                    {"strategy_" + k: v for k, v in strategy.items() if k != "params"}
                )
                flat.update({key: strategy["params"].get(key) for key in PARAM_KEYS})
                flat["strategy_params"] = {
                    k: v for k, v in strategy["params"].items() if k not in PARAM_KEYS
                }
                ps.append(
                    [
                        json.dumps(flat.get(k), ensure_ascii=False)
                        if isinstance(flat.get(k), (dict, list))
                        else flat.get(k)
                        for k in POINT_COLUMNS.values()
                    ]
                )
    else:
        ws = wb.create_sheet("数据")
        rows = payload.get("rows", [])
        headers = list(rows[0]) if rows else ["time", "value"]
        ws.append(headers)
        for row in rows:
            ws.append(
                [
                    json.dumps(row.get(k), ensure_ascii=False)
                    if isinstance(row.get(k), (dict, list))
                    else row.get(k)
                    for k in headers
                ]
            )
    for ws in wb:
        text_cells(ws)
        freeze_header(ws)
        ws.auto_filter.ref = ws.dimensions
    wb.save(destination)
    return {"path": str(destination)}


def parse_workbook(source, mode, current=None):
    errors = []

    def json_field(value, sheet, row, field, fallback):
        try:
            return json.loads(value)
        except (ValueError, TypeError) as exc:
            errors.append(
                {
                    "sheet": sheet,
                    "row": row,
                    "field": field,
                    "message": "JSON 格式无效：" + str(exc)[:200],
                }
            )
            return fallback

    if Path(source).stat().st_size > 20 * 1024 * 1024:
        return {
            "errors": [
                {
                    "sheet": "文件",
                    "row": 0,
                    "field": "大小",
                    "message": "文件超过 20MiB",
                }
            ]
        }
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 10000
            or sum(i.file_size for i in entries) > 100 * 1024 * 1024
        ):
            return {
                "errors": [
                    {
                        "sheet": "文件",
                        "row": 0,
                        "field": "解压大小",
                        "message": "压缩包解压超过限制",
                    }
                ]
            }
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        if wb.sheetnames == ["点位"]:
            return parse_points(wb, current)
        if not {"格式", "设备", "点位"}.issubset(wb.sheetnames):
            raise ValueError("缺少格式、设备或点位工作表")
        meta = {
            row[0]: row[1]
            for row in wb["格式"].iter_rows(
                min_row=1, max_row=10, max_col=2, values_only=True
            )
            if row[0]
        }
        if (
            meta.get("格式版本") != 1
            or meta.get("类型") != "config"
            or meta.get("地址格式") != "protocol_zero"
        ):
            raise ValueError("仅支持版本 1 的配置文件和从 0 开始的协议地址")
        devices, locations = [], {}

        def rows(sheet, columns, maximum):
            ws = wb[sheet]
            if ws.max_column > len(columns) or ws.max_row > maximum + 1:
                raise ValueError(f"{sheet} 行／列数超过限制")
            iterator = ws.iter_rows(values_only=True)
            headers = next(iterator)
            if len(set(headers)) != len(headers) or any(
                h not in columns for h in headers
            ):
                raise ValueError(f"{sheet} 列名非法或重复")
            for index, values in enumerate(iterator, 2):
                if index > maximum + 1:
                    raise ValueError(f"{sheet} 超过行数限制")
                if all(v is None for v in values):
                    continue
                yield (
                    index,
                    {columns[h]: v for h, v in zip(headers, values) if v is not None},
                )

        for row, item in rows("设备", DEVICE_COLUMNS, 16):
            if "id" not in item:
                if mode != "new":
                    errors.append(
                        {
                            "sheet": "设备",
                            "row": row,
                            "field": "设备 ID",
                            "message": "缺少稳定 ID，只能使用新增模式",
                        }
                    )
                    continue
                item["id"] = uid()
            for key in ("functions", "faults", "valid_ranges", "identity"):
                if key in item:
                    field = next(h for h, k in DEVICE_COLUMNS.items() if k == key)
                    item[key] = json_field(item[key], "设备", row, field, None)
            item["points"] = []
            locations[("devices", len(devices))] = ("设备", row)
            devices.append(item)
        by_id = {}
        for index, d in enumerate(devices):
            if d["id"] in by_id:
                errors.append(
                    {
                        "sheet": "设备",
                        "row": locations[("devices", index)][1],
                        "field": "设备 ID",
                        "message": "设备 ID 重复",
                    }
                )
            by_id[d["id"]] = d
        point_ids = set()
        for row, item in rows("点位", POINT_COLUMNS, 10000):
            device_id = item.pop("device_id", None)
            if device_id is None and mode == "new" and len(devices) == 1:
                device_id = devices[0]["id"]
            if device_id not in by_id:
                errors.append(
                    {
                        "sheet": "点位",
                        "row": row,
                        "field": "设备 ID",
                        "message": "设备引用不存在",
                    }
                )
                continue
            if "id" not in item:
                if mode != "new":
                    errors.append(
                        {
                            "sheet": "点位",
                            "row": row,
                            "field": "点位 ID",
                            "message": "缺少稳定 ID，只能使用新增模式",
                        }
                    )
                    continue
                item["id"] = uid()
            if item["id"] in point_ids:
                errors.append(
                    {
                        "sheet": "点位",
                        "row": row,
                        "field": "点位 ID",
                        "message": "点位 ID 重复",
                    }
                )
            point_ids.add(item["id"])
            params = json_field(
                item.pop("strategy_params", "{}"), "点位", row, "附加策略参数", {}
            )
            if not isinstance(params, dict):
                errors.append(
                    {
                        "sheet": "点位",
                        "row": row,
                        "field": "附加策略参数",
                        "message": "必须是 JSON 对象",
                    }
                )
                params = {}
            for key in PARAM_KEYS:
                if key in item:
                    params[key] = item.pop(key)
            strategy = {
                k[len("strategy_") :]: item.pop(k)
                for k in list(item)
                if k.startswith("strategy_")
            }
            if "dependencies" in strategy:
                strategy["dependencies"] = json_field(
                    strategy["dependencies"], "点位", row, "依赖点位", []
                )
            strategy["params"] = params
            item["strategy"] = strategy
            device = by_id[device_id]
            d_index = devices.index(device)
            locations[("devices", d_index, "points", len(device["points"]))] = (
                "点位",
                row,
            )
            device["points"].append(item)
        if errors:
            return {"errors": errors}
        payload = {"devices": devices, "settings": json.loads(meta.get("设置", "{}"))}
        try:
            validated = Configuration.model_validate(payload)
        except ValidationError as exc:
            for error in exc.errors(include_input=False, include_url=False):
                loc = error["loc"]
                sheet, row = locations.get(loc[:4], locations.get(loc[:2], ("格式", 0)))
                errors.append(
                    {
                        "sheet": sheet,
                        "row": row,
                        "field": str(loc[-1]) if loc else "配置",
                        "message": error["msg"],
                    }
                )
            return {"errors": errors}
        return {"errors": [], "config": validated.model_dump()}
    finally:
        wb.close()


def child(operation, payload, destination, result):
    try:
        data = (
            parse_workbook(payload["path"], payload["mode"], payload.get("current"))
            if operation == "parse"
            else write_workbook(payload, destination)
        )
        Path(result).write_text(
            json.dumps(data, ensure_ascii=False, allow_nan=False), encoding="utf-8"
        )
    except Exception as exc:
        Path(result).write_text(
            json.dumps(
                {
                    "errors": [
                        {
                            "sheet": "文件",
                            "row": 0,
                            "field": "格式",
                            "message": str(exc)[:512],
                        }
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )


class HeavyTasks:
    def __init__(self, directory, budget=256 * 1024 * 1024, timeout=60):
        self.directory, self.budget, self.timeout = Path(directory), budget, timeout
        self.semaphore = asyncio.Semaphore(1)
        self.waiting = 0
        self.active = set()
        self.jobs = set()
        self.closed = False

    async def run(self, operation, payload, destination=None):
        if self.closed:
            raise DomainError("重任务服务已停止", 503)
        if self.waiting >= 3:
            raise DomainError("重任务队列已满", 429)
        self.waiting += 1
        task = asyncio.current_task()
        self.jobs.add(task)
        token = uid()
        result = self.directory / f"{token}.json"
        destination = Path(destination or self.directory / f"{token}.xlsx")
        process = None
        exported = False
        try:
            async with self.semaphore:
                used = directory_bytes(self.directory)
                if used >= self.budget:
                    raise DomainError("临时文件预算不足", 507)
                process = multiprocessing.get_context("spawn").Process(
                    target=child, args=(operation, payload, destination, result)
                )
                process.start()
                self.active.add(process)
                deadline = time.monotonic() + self.timeout
                while process.is_alive():
                    if time.monotonic() > deadline:
                        raise DomainError("任务超时，请缩小范围或检查文件", 408)
                    if directory_bytes(self.directory) > self.budget:
                        raise DomainError("临时文件超过容量上限", 507)
                    await asyncio.sleep(0.05)
                process.join(timeout=0)
                if directory_bytes(self.directory) > self.budget:
                    raise DomainError("临时文件超过容量上限", 507)
                if not result.exists():
                    raise DomainError("重任务进程异常退出", 503)
                if result.stat().st_size > 20 * 1024 * 1024:
                    raise DomainError("任务结果过大", 413)
                data = json.loads(result.read_text(encoding="utf-8"))
                exported = operation == "export" and bool(data.get("path"))
                return data
        finally:
            self.waiting -= 1
            if process:
                if process.is_alive():
                    process.terminate()
                    await asyncio.to_thread(process.join, 1)
                    if process.is_alive():
                        process.kill()
                        await asyncio.to_thread(process.join, 1)
                self.active.discard(process)
                process.close()
            result.unlink(missing_ok=True)
            if not exported:
                destination.unlink(missing_ok=True)
            if operation == "parse":
                Path(payload["path"]).unlink(missing_ok=True)
            self.jobs.discard(task)

    async def close(self):
        self.closed = True
        jobs = list(self.jobs)
        for task in jobs:
            task.cancel()
        # Each job owns termination, bounded join and file cleanup. Waiting
        # here prevents queued jobs from spawning workers after shutdown.
        await asyncio.gather(*jobs, return_exceptions=True)
