"""Human-readable Modbus address map, generated in the Excel worker."""

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from .codec import width
from .point_excel import AREAS, freeze_header

HEADERS = [
    "设备名称",
    "设备 ID",
    "监听 IP",
    "端口",
    "Unit ID",
    "点位名称",
    "分组",
    "数据区",
    "读功能码",
    "写功能码",
    "协议地址（从0）",
    "参考编号（六位）",
    "占用长度",
    "存储单位",
    "数据类型",
    "字节序",
    "字序",
    "倍率",
    "偏移",
    "工程单位",
    "主机可写",
    "初始值",
    "说明",
]
READ_CODES = {"coil": (1,), "discrete": (2,), "holding": (3, 23), "input": (4,)}
WRITE_CODES = {"coil": (5, 15), "discrete": (), "holding": (6, 16, 22, 23), "input": ()}
PREFIX = {"coil": "0", "discrete": "1", "holding": "4", "input": "3"}


def write_point_table(wb, devices):
    ws = wb.active
    ws.title = "Modbus点表"
    ws.append(HEADERS)
    areas = {value: name for name, value in AREAS.items()}
    order = {area: index for index, area in enumerate(areas)}
    for device in devices:
        functions = set(device["functions"])
        for point in sorted(
            device["points"], key=lambda p: (order[p["area"]], p["address"])
        ):
            area = point["area"]
            codes = lambda allowed: (
                " / ".join(f"{code:02d}" for code in allowed if code in functions)
                or "—"
            )
            writable = point["writable"] and any(
                code in functions for code in WRITE_CODES[area]
            )
            length = width(point["type"])
            ws.append(
                [
                    device["name"],
                    device["id"],
                    device["host"],
                    device["port"],
                    device["unit_id"],
                    point["name"],
                    point["group"],
                    areas[area],
                    codes(READ_CODES[area]),
                    codes(WRITE_CODES[area]) if writable else "—",
                    point["address"],
                    PREFIX[area] + f"{point['address'] + 1:05d}",
                    length,
                    "位" if point["type"] == "Bool" else "寄存器（16位）",
                    point["type"],
                    ("大端" if point["byte_order"] == "big" else "小端")
                    if point["type"] != "Bool"
                    else "—",
                    ("高字在前" if point["word_order"] == "big" else "低字在前")
                    if length > 1
                    else "—",
                    point["scale"],
                    point["offset"],
                    point["unit"],
                    "是" if writable else "否",
                    point["initial"],
                    point["description"],
                ]
            )
    notes = {
        "监听 IP": "0.0.0.0/:: 表示监听所有本机地址；远程主机应连接服务所在电脑实际 IP。",
        "读功能码": "仅列出设备启用的功能码；23 为读写多个保持寄存器，调用时还需合法可写地址。",
        "写功能码": "仅列出设备启用且点位允许写入的功能码；22 为掩码写寄存器，23 为读写多个寄存器。",
        "协议地址（从0）": "Modbus 请求使用的零基起始地址。多寄存器点位按占用长度连续占用地址。",
        "参考编号（六位）": "对照编号，不是请求地址：0=线圈，1=离散输入，3=输入寄存器，4=保持寄存器；后五位=协议地址+1。例如保持寄存器地址0为400001（常见五位写法40001）。",
        "倍率": "工程值 = 解码原始值 × 倍率 + 偏移。",
        "初始值": "配置中的工程初始值；实时值请使用导出当前快照。本点表供协议对接，不用于配置导入。",
    }
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245C86")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = (
            24 if cell.value in ("设备 ID", "说明") else 20
        )
        if cell.value in notes:
            cell.comment = Comment(notes[cell.value], "Modbus Simulator")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                # Avoid formula injection and retain leading zeros in references.
                cell.data_type = "s"
                cell.number_format = "@"
    ws.row_dimensions[1].height = 26
    freeze_header(ws)
    ws.auto_filter.ref = ws.dimensions
