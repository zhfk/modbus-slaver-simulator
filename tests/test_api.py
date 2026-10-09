import asyncio
import json
import threading
import time

import pytest
from fastapi.testclient import TestClient
from simulator.api import TemporaryFileResponse, create_app
from simulator.excel import parse_workbook, write_workbook
from simulator.models import Configuration, thermal_template


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        yield client


@pytest.mark.parametrize("source", ["default", "environment", "argument"])
def test_health_reports_effective_absolute_data_directory(
    tmp_path, monkeypatch, source
):
    monkeypatch.chdir(tmp_path)
    default = tmp_path / "默认数据 # 目录"
    monkeypatch.setattr("simulator.api.default_data_dir", lambda: default)
    argument = None
    if source == "default":
        monkeypatch.delenv("MODBUS_DATA_DIR", raising=False)
        expected = default
    else:
        monkeypatch.setenv("MODBUS_DATA_DIR", "环境数据 # 目录")
        expected = tmp_path / "环境数据 # 目录"
        if source == "argument":
            argument = "指定数据 # 目录"
            expected = tmp_path / argument
    with TestClient(create_app(argument)) as http:
        assert http.get("/api/health").json()["storage"]["data_dir"] == str(
            expected.resolve()
        )
        assert (expected / "config.db").is_file()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "failure", ["disconnect", "cancel", "invalid-range", "success"]
)
async def test_export_response_cleans_file_even_when_download_is_interrupted(
    tmp_path, failure
):
    path = tmp_path / "export.xlsx"
    path.write_bytes(b"export content")
    response = TemporaryFileResponse(path, filename="export.xlsx")
    sent = []

    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message):
        sent.append(message)
        if message["type"] == "http.response.body":
            if failure == "disconnect":
                raise OSError("client disconnected")
            if failure == "cancel":
                raise asyncio.CancelledError()

    scope = {
        "type": "http",
        "method": "GET",
        "headers": [(b"range", b"invalid")] if failure == "invalid-range" else [],
    }
    if failure in ("disconnect", "cancel"):
        with pytest.raises(
            OSError if failure == "disconnect" else asyncio.CancelledError
        ):
            await response(scope, receive, send)
    else:
        await response(scope, receive, send)
        assert sent[0]["status"] == (400 if failure == "invalid-range" else 200)
    assert not path.exists()


def template(client):
    result = client.post("/api/templates/thermal", json={"version": 0, "name": "温控"})
    assert result.status_code == 200, result.text
    return result.json()["id"]


def test_api_state_and_validation(client):
    key = template(client)
    config = client.get("/api/config").json()
    assert config["version"] == 1
    items = client.get(f"/api/devices/{key}/points").json()["items"]
    point = items[1]
    preview = client.post(
        f"/api/devices/{key}/preview-value", json={"id": point["id"], "value": 25.35}
    )
    assert preview.json()["raw"] == [254]
    response = client.post(
        f"/api/devices/{key}/assign", json={"items": [{"id": point["id"], "value": 80}]}
    )
    assert response.status_code == 200
    assert response.json()["items"][0]["value"] == 80
    assert client.get("/api/config").json()["devices"][0]["points"][1]["initial"] == 60
    bad = json.loads(json.dumps(config))
    bad["devices"][0]["points"][1]["scale"] = 0
    assert client.put("/api/config", json=bad).status_code == 422
    assert client.get("/api/config").json()["version"] == 1
    assert client.get("/api/health/live").status_code == 200
    assert client.get("/api/health/ready").status_code == 200
    assert client.get("/api/not-an-api").status_code == 404
    assert client.get("/assets/missing.js").status_code == 404
    assert client.get(f"/api/devices/{key}/points?size=10000").status_code == 422


@pytest.mark.parametrize("invalid_id", [[], {}, None, 1])
def test_invalid_point_id_returns_client_error_and_keeps_state(client, invalid_id):
    key = template(client)
    before = client.get(f"/api/devices/{key}/points").json()
    for action, payload in (
        ("assign", {"items": [{"id": invalid_id, "value": 80}]}),
        ("preview-value", {"id": invalid_id, "value": 80}),
    ):
        assert (
            client.post(f"/api/devices/{key}/{action}", json=payload).status_code == 422
        )
    assert client.get(f"/api/devices/{key}/points").json()["items"] == before["items"]
    assert client.get("/api/health/ready").status_code == 200


@pytest.mark.parametrize(
    "kind,params",
    [
        ("fixed", {"value": None}),
        ("sine", {"period": []}),
        ("expression", {"expression": "(("}),
        ("sequence", {"values": 1}),
        ("replay", {"values": [[0, None]]}),
    ],
)
def test_invalid_strategy_is_validation_error_without_config_commit(
    client, kind, params
):
    template(client)
    config = client.get("/api/config").json()
    bad = json.loads(json.dumps(config))
    bad["devices"][0]["points"][2]["strategy"].update(
        kind=kind, params=params, dependencies=[]
    )
    assert client.put("/api/config", json=bad).status_code == 422
    assert client.get("/api/config").json() == config


def test_large_assignment_does_not_change_previous_value(client):
    key = template(client)
    point = client.get(f"/api/devices/{key}/points").json()["items"][1]
    for action, payload in (
        ("assign", {"items": [{"id": point["id"], "value": 1e100}]}),
        ("preview-value", {"id": point["id"], "value": 1e100}),
    ):
        assert (
            client.post(f"/api/devices/{key}/{action}", json=payload).status_code == 422
        )
    assert (
        client.get(f"/api/devices/{key}/points").json()["items"][1]["value"]
        == point["value"]
    )


@pytest.mark.parametrize("port", [None, [], 1502.5])
def test_invalid_template_port_does_not_create_device(client, port):
    assert (
        client.post(
            "/api/templates/thermal", json={"version": 0, "port": port}
        ).status_code
        == 422
    )
    assert client.get("/api/config").json()["devices"] == []


def test_disabled_history_still_exports_existing_samples(client):
    from openpyxl import load_workbook
    from io import BytesIO

    key = template(client)
    point = client.get("/api/config").json()["devices"][0]["points"][1]
    storage = client.app.state.context["storage"]
    now = time.time()
    client.portal.call(
        storage.telemetry_db.call,
        lambda c: (
            c.execute(
                "INSERT INTO samples(time,device,point,value,raw,metadata,version) VALUES (?,?,?,?,?,?,?)",
                (now, key, point["id"], 60, "[600]", '{"unit":"℃"}', 1),
            ),
            c.commit(),
        ),
    )
    assert not client.get("/api/config").json()["settings"]["history_enabled"]
    response = client.get(
        f"/api/export/history?device={key}&start={now - 1}&end={now + 1}"
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    try:
        rows = list(workbook["数据"].values)
        assert len(rows) == 2 and rows[1][rows[0].index("value")] == 60
    finally:
        workbook.close()


@pytest.mark.parametrize("query", [[], {"ids": [[]]}, {"ids": [None]}])
def test_invalid_live_subscription_closes_cleanly(client, query):
    from starlette.websockets import WebSocketDisconnect

    key = template(client)
    if isinstance(query, dict):
        query = {**query, "device": key}
    with client.websocket_connect("/api/live") as ws:
        ws.send_json(query)
        with pytest.raises(WebSocketDisconnect) as error:
            ws.receive_json()
        assert error.value.code == 1008
    assert client.app.state.context["sessions"] == 0


def test_excel_roundtrip_formula_and_invalid_row(tmp_path):
    device = thermal_template(name="=DANGEROUS()")
    device.points[0].description = "=1+1"
    config = Configuration(devices=[device])
    path = tmp_path / "config.xlsx"
    write_workbook({**config.model_dump(), "kind": "config"}, path)
    result = parse_workbook(path, "update")
    assert result["errors"] == []
    restored = Configuration.model_validate(result["config"])
    assert restored.devices[0].model_dump() == device.model_dump()
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb["设备"]["B2"].data_type == "s"
    header = [c.value for c in wb["点位"][1]]
    wb["点位"].cell(3, header.index("倍率") + 1, 0)
    wb.save(path)
    errors = parse_workbook(path, "update")["errors"]
    assert errors[0]["sheet"] == "点位" and errors[0]["row"] == 3


def test_actual_excel_worker_export_import(client):
    key = template(client)
    export = client.get(f"/api/export/config?device={key}")
    assert export.status_code == 200, export.text
    assert export.content[:2] == b"PK"
    preview = client.post(
        "/api/import/preview?mode=update",
        files={
            "file": (
                "points.xlsx",
                export.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert not data["errors"] and data["changed"] == 4 and data["deleted"] == 0
    assert client.post("/api/import/apply", json=data["config"]).status_code == 200
    assert client.post("/api/import/apply", json=data["config"]).status_code == 409
    snapshot = client.get(f"/api/export/snapshot?device={key}")
    invalid = client.post(
        "/api/import/preview", files={"file": ("snapshot.xlsx", snapshot.content)}
    )
    assert invalid.json()["errors"]


def test_excel_json_and_duplicate_id_errors_identify_exact_source_row(tmp_path):
    from openpyxl import load_workbook

    d = thermal_template()
    path = tmp_path / "invalid.xlsx"
    write_workbook({**Configuration(devices=[d]).model_dump(), "kind": "config"}, path)
    wb = load_workbook(path)
    from simulator.excel import POINT_COLUMNS

    headers = list(POINT_COLUMNS.values())
    wb["点位"].cell(3, headers.index("strategy_params") + 1, "{broken")
    wb["点位"].cell(4, headers.index("id") + 1, d.points[0].id)
    wb.save(path)
    errors = parse_workbook(path, "update")["errors"]
    assert any(
        e["sheet"] == "点位" and e["row"] == 3 and "JSON" in e["message"]
        for e in errors
    )
    assert any(
        e["sheet"] == "点位" and e["row"] == 4 and "重复" in e["message"]
        for e in errors
    )


def test_websocket_complete_snapshot_and_size_limit(client):
    key = template(client)
    config = client.get("/api/config").json()
    p = config["devices"][0]["points"][0]
    with client.websocket_connect("/api/live") as ws:
        ws.send_json({"device": key, "ids": [p["id"]]})
        frame = ws.receive_json()
        assert frame["complete"] and frame["items"][0]["id"] == p["id"]
        assert frame["device"]["id"] == key


def test_health_detects_stalled_database_job_even_with_empty_queue(client):
    db = client.app.state.context["storage"].config_db
    entered, release = threading.Event(), threading.Event()

    def blocked(conn):
        entered.set()
        assert release.wait(5)

    future = db.submit(blocked)
    assert entered.wait(1)
    try:
        assert db.queue.empty() and db.busy
        # Test the deadline without making the test wait fifteen seconds.
        db.busy_since = time.monotonic() - 16
        response = client.get("/api/health/live")
        assert response.status_code == 503
        assert response.json()["runtime"] and not response.json()["storage"]
    finally:
        release.set()
        future.result(timeout=2)
    assert client.get("/api/health/live").status_code == 200


def test_corrupt_history_preserves_file_and_does_not_block_autostart_or_config(
    tmp_path,
):
    import socket
    from pymodbus.client import ModbusTcpClient

    with socket.socket() as available:
        available.bind(("127.0.0.1", 0))
        port = available.getsockname()[1]
    with TestClient(create_app(tmp_path)) as http:
        d = thermal_template(port=port)
        d.auto_start = True
        assert (
            http.put(
                "/api/config", json=Configuration(devices=[d]).model_dump()
            ).status_code
            == 200
        )
    damaged = b"damaged telemetry database; preserve evidence"
    (tmp_path / "telemetry.db").write_bytes(damaged)
    with TestClient(create_app(tmp_path)) as http:
        assert http.get("/api/health/ready").status_code == 200
        metric = http.get("/api/health").json()
        assert (
            not metric["storage"]["error"]
            and "历史库" in metric["storage"]["history_error"]
        )
        assert (tmp_path / "telemetry.db").read_bytes() == damaged
        protocol = ModbusTcpClient("127.0.0.1", port=port, timeout=2)
        try:
            assert protocol.connect()
            assert protocol.read_holding_registers(
                0, count=1, device_id=1
            ).registers == [600]
        finally:
            protocol.close()
        config = http.get("/api/config").json()
        config["devices"][0]["name"] = "history fault isolated"
        assert http.put("/api/config", json=config).status_code == 200


def test_application_log_rotates_and_keeps_total_below_fifty_mib(client):
    import logging

    logger = logging.getLogger("simulator")
    handler = next(h for h in logger.handlers if getattr(h, "maxBytes", 0))
    directory = __import__("pathlib").Path(handler.baseFilename).parent
    for i in range(120):
        logger.info("rotation sample %s %s", i, "x" * (512 * 1024))
    handler.flush()
    files = list(directory.glob("application.log*"))
    assert len(files) == 5
    assert sum(p.stat().st_size for p in files) <= 50 * 1024**2
    assert "rotation sample 119" in (directory / "application.log").read_text()
