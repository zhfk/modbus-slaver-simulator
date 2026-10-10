from io import BytesIO

import pytest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from simulator.api import create_app
from simulator.excel import parse_workbook, write_workbook
from simulator.models import Configuration, Device, Point
from simulator.point_table import HEADERS


def test_protocol_map_addresses_encoding_permissions_and_text_safety(tmp_path):
    d = Device(
        id="000012345678901234567890",
        name="=device",
        functions=[1, 2, 3, 4, 6],
        points=[
            Point(
                id="c",
                name="coil",
                area="coil",
                type="Bool",
                address=65535,
                initial=True,
                writable=True,
            ),
            Point(id="d", name="discrete", area="discrete", type="Bool", address=0),
            Point(
                id="h",
                name="=SUM(A1)",
                type="Float32",
                address=0,
                byte_order="little",
                word_order="little",
                scale=0.1,
                offset=2,
                unit="℃",
                writable=True,
            ),
            Point(id="r", name="readonly", address=4),
            Point(id="i", name="input", area="input", type="Float64", address=0),
        ],
    )
    path = tmp_path / "table.xlsx"
    write_workbook({"kind": "point-table", "devices": [d.model_dump()]}, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Modbus点表"]
    ws = wb.active
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_column == 20
    assert not {"设备 ID", "分组", "初始值"}.intersection(HEADERS)
    rows = {
        row[4]: dict(zip(HEADERS, row))
        for row in ws.iter_rows(min_row=2, values_only=True)
    }
    assert rows["coil"]["参考编号（六位）"] == "065536"
    assert rows["coil"]["读功能码"] == "01" and rows["coil"]["写功能码"] == "—"
    assert rows["coil"]["主机可写"] == "否"
    assert rows["discrete"]["参考编号（六位）"] == "100001"
    h = rows["=SUM(A1)"]
    assert h["协议地址（从0）"] == 0 and h["参考编号（六位）"] == "400001"
    assert (h["占用长度"], h["存储单位"], h["字节序"], h["字序"]) == (
        2,
        "寄存器（16位）",
        "小端",
        "低字在前",
    )
    assert (h["倍率"], h["偏移"], h["工程单位"], h["写功能码"]) == (0.1, 2, "℃", "06")
    assert rows["readonly"]["写功能码"] == "—" and rows["input"]["占用长度"] == 4
    assert rows["input"]["参考编号（六位）"] == "300001"
    assert ws["E4"].data_type == "s" and ws["J4"].number_format == "@"
    assert ws.freeze_panes is None and ws.sheet_view.pane is None
    assert ws.sheet_view.selection[0].activeCell == "A1"
    assert ws["I1"].comment and ws["J1"].comment
    wb.close()
    with pytest.raises(ValueError, match="缺少"):
        parse_workbook(path, "update")


def test_actual_point_map_worker_selection_shared_endpoint_and_backup(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        config = Configuration(
            devices=[
                Device(
                    id="one",
                    name="一号",
                    points=[
                        Point(id="p", name="温度"),
                        Point(id="other", name="其他", address=2),
                    ],
                ),
                Device(
                    id="two",
                    name="二号",
                    unit_id=2,
                    points=[Point(id="q", name="温度")],
                ),
            ]
        ).model_dump()
        assert client.put("/api/config", json=config).status_code == 200
        before = client.get("/api/config").json()
        for query, count, units in [("", 3, {1, 2}), ("?device=one&ids=p", 1, {1})]:
            result = client.get("/api/export/point-table" + query)
            assert result.status_code == 200, result.text
            assert "modbus-point-table.xlsx" in result.headers["content-disposition"]
            wb = load_workbook(BytesIO(result.content))
            assert wb.sheetnames == ["Modbus点表"] and wb.active.max_row == count + 1
            assert wb.active.max_column == 20
            assert not {"设备 ID", "分组", "初始值"}.intersection(
                cell.value for cell in wb.active[1]
            )
            assert {
                r[3] for r in wb.active.iter_rows(min_row=2, values_only=True)
            } == units
            wb.close()
        assert client.get("/api/export/point-table?ids=missing").status_code == 422
        assert client.get("/api/export/point-table?device=missing").status_code == 404
        backup = load_workbook(BytesIO(client.get("/api/export/config").content))
        assert backup.sheetnames == ["格式", "设备", "点位"]
        backup.close()
        assert client.get("/api/config").json() == before
