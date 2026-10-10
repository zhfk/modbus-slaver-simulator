from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from simulator.api import create_app
from simulator.codec import decode, encode
from simulator.excel import parse_workbook, write_workbook
from simulator.models import Configuration, Device, Point, Strategy
from simulator.point_excel import COLUMNS
from simulator.runtime import DeviceRuntime


def workbook(tmp_path, rows):
    path = tmp_path / "points.xlsx"
    write_workbook({"kind": "template"}, path)
    wb = load_workbook(path)
    for row in rows:
        wb["点位"].append([row.get(column) for column in COLUMNS])
    wb.save(path)
    wb.close()
    return path


def row(device="device", name="温度", **changes):
    return {
        "设备 ID": device,
        "名称": name,
        "数据区": "保持寄存器",
        "协议地址": 0,
        "类型": "Float32",
        **changes,
    }


def test_template_has_only_fourteen_columns_and_native_dropdowns(tmp_path):
    path = workbook(tmp_path, [])
    wb = load_workbook(path)
    assert wb.sheetnames == ["点位"]
    ws = wb["点位"]
    assert [c.value for c in ws[1]] == list(COLUMNS)
    assert ws.max_row == 1 and ws.max_column == 14
    rules = list(ws.data_validations.dataValidation)
    assert len(rules) == 4 and all(
        r.type == "list" and r.showErrorMessage for r in rules
    )
    assert any("线圈,离散输入,保持寄存器,输入寄存器" in r.formula1 for r in rules)
    assert any("均匀随机" in r.formula1 for r in rules)
    assert ws["A1"].comment and ws["J1"].comment
    wb.close()


def test_upsert_uses_device_and_exact_name_retains_ids_references_and_other_points(
    tmp_path,
):
    first = Device(
        id="device",
        points=[
            Point(id="old", name="温度", type="Float32", initial=12),
            Point(id="keep", name="未导入", address=4),
            Point(
                id="dep",
                name="依赖",
                address=5,
                strategy=Strategy(
                    kind="expression", dependencies=["old"], params={"expression": "x0"}
                ),
            ),
        ],
    )
    second = Device(
        id="other",
        unit_id=2,
        points=[Point(id="second", name="温度", type="Float32", initial=8)],
    )
    current = Configuration(version=7, devices=[first, second]).model_dump()
    path = workbook(
        tmp_path,
        [
            row(**{"分组": "车间", "初始值": 20, "最大值": 40}),
            row(name="新点位", **{"协议地址": 8}),
            row("other", **{"初始值": 15}),
        ],
    )
    result = parse_workbook(path, "update", current)
    assert result["errors"] == [] and result["format"] == "points"
    assert (result["added"], result["changed"], result["deleted"]) == (1, 2, 0)
    config = result["config"]
    assert config["version"] == 7
    a, b = config["devices"]
    assert a["points"][0]["id"] == "old" and a["points"][0]["group"] == "车间"
    assert a["points"][1] == first.points[1].model_dump()
    assert a["points"][2]["strategy"]["dependencies"] == ["old"]
    assert a["points"][3]["id"] not in {"old", "keep", "dep", "second"}
    assert b["points"][0]["id"] == "second" and b["points"][0]["initial"] == 15
    assert current["devices"][0]["points"][0]["initial"] == 12
    for mode in ("new", "replace"):
        assert parse_workbook(path, mode, current)["deleted"] == 0


@pytest.mark.parametrize(
    "kind", ["Bool", "Int16", "UInt16", "Int32", "UInt32", "Float32", "Float64"]
)
def test_defaults_and_random_initial_are_encodable_for_every_type(tmp_path, kind):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    path = workbook(
        tmp_path,
        [row(**{"类型": kind, "数据区": "线圈" if kind == "Bool" else "保持寄存器"})],
    )
    result = parse_workbook(path, "update", current)
    assert not result["errors"], result
    p = Point.model_validate(result["config"]["devices"][0]["points"][0])
    assert 0 <= p.initial <= (1 if kind == "Bool" else 100)
    assert p.byte_order == p.word_order == "big" and p.write_mode == "hold"
    assert (
        p.scale == 1
        and p.offset == 0
        and p.precision == (2 if kind.startswith("Float") else 0)
    )
    assert p.strategy.kind == "random" and p.strategy.enabled
    assert p.strategy.interval == p.strategy.seed == 1
    assert abs(decode(p, encode(p, p.initial)) - p.initial) < 0.001


@pytest.mark.parametrize(
    "strategy", ["无策略", "固定值", "均匀随机", "正弦波", "斜坡", "随机游走"]
)
def test_dropdown_strategies_generate_in_the_configured_range(tmp_path, strategy):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    result = parse_workbook(
        workbook(tmp_path, [row(**{"策略类型": strategy, "最小值": 20, "最大值": 30})]),
        "update",
        current,
    )
    assert not result["errors"], result
    device = Device.model_validate(result["config"]["devices"][0])
    runtime = DeviceRuntime(device)
    point = device.points[0]
    for time in range(61):
        runtime.states[point.id].elapsed = time
        value = runtime.strategy_value(point.id, {point.id: point.initial}, 1)
        assert 20 <= value <= 30
        encode(point, value)


@pytest.mark.parametrize("area", ["离散输入", "输入寄存器"])
def test_input_areas_cannot_be_writable(tmp_path, area):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    r = row(
        **{
            "数据区": area,
            "类型": "Bool" if area == "离散输入" else "UInt16",
            "主机可写": "是",
        }
    )
    result = parse_workbook(workbook(tmp_path, [r]), "update", current)
    assert any(e["row"] == 2 and "只读" in e["message"] for e in result["errors"])


@pytest.mark.parametrize(
    "changes, message",
    [
        ({"设备 ID": "unknown"}, "设备 ID"),
        ({"名称": ""}, "名称"),
        ({"倍率": 0}, "倍率"),
        ({"最大值": 2, "类型": "Bool", "数据区": "线圈"}, "最小值"),
        ({"数据区": "holding"}, "数据区"),
        ({"最小值": 80, "最大值": 20}, "顺序"),
        ({"策略类型": "温控联动"}, "复杂策略"),
        ({"类型": "Bad"}, "类型"),
        ({"初始值": 100000, "类型": "UInt16"}, "范围"),
    ],
)
def test_bad_rows_identify_source_and_do_not_return_a_candidate(
    tmp_path, changes, message
):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    result = parse_workbook(workbook(tmp_path, [row(**changes)]), "update", current)
    assert result["errors"] and "config" not in result
    assert result["errors"][0]["row"] == 2
    assert message in result["errors"][0]["message"]
    assert current["devices"][0]["points"] == []


def test_duplicate_names_and_overlapping_addresses_fail_atomically(tmp_path):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    for rows, message in [
        ([row(), row()], "重复"),
        ([row(), row(name="另一点位", **{"协议地址": 1})], "重叠"),
    ]:
        result = parse_workbook(workbook(tmp_path, rows), "update", current)
        assert "config" not in result and any(
            message in e["message"] for e in result["errors"]
        )
    d = Device(
        id="device",
        points=[Point(name="温度", address=0), Point(name="温度", address=1)],
    )
    result = parse_workbook(
        workbook(tmp_path, [row()]), "update", Configuration(devices=[d]).model_dump()
    )
    assert "多个同名" in result["errors"][0]["message"]


def test_scaled_integer_random_initial_and_boolean_random_are_valid(tmp_path):
    current = Configuration(devices=[Device(id="device")]).model_dump()
    result = parse_workbook(
        workbook(
            tmp_path,
            [
                row(
                    **{
                        "类型": "Int16",
                        "倍率": -0.1,
                        "偏移": 50,
                        "最小值": 40,
                        "最大值": 60,
                    }
                ),
                row(name="开关", **{"数据区": "线圈", "类型": "Bool", "协议地址": 0}),
            ],
        ),
        "update",
        current,
    )
    assert not result["errors"], result
    device = Device.model_validate(result["config"]["devices"][0])
    p, b = device.points
    assert 40 <= p.initial <= 60 and p.precision == 0
    assert decode(p, encode(p, p.initial)) == pytest.approx(p.initial)
    runtime = DeviceRuntime(device)
    values = {b.id: b.initial}
    samples = {runtime.strategy_value(b.id, values, 1) for _ in range(100)}
    assert samples == {0, 1}
    assert all(encode(b, value) in ([0], [1]) for value in samples)
    with pytest.raises(ValueError, match="Bool 均匀随机"):
        Point(
            name="bad",
            area="coil",
            type="Bool",
            strategy=Strategy(kind="random", params={"min": 0, "max": 2}),
        )


def test_actual_worker_download_upsert_apply_preserves_random_preview_and_version(
    tmp_path,
):
    with TestClient(create_app(tmp_path / "data")) as client:
        config = Configuration(
            devices=[
                Device(id="device", points=[Point(id="old", name="旧点位", address=4)])
            ]
        ).model_dump()
        assert client.put("/api/config", json=config).status_code == 200
        response = client.get("/api/export/template")
        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        assert wb.sheetnames == ["点位"]
        for r in [
            row(name="旧点位", **{"协议地址": 4, "类型": "UInt16", "初始值": 12}),
            row(),
        ]:
            wb["点位"].append([r.get(c) for c in COLUMNS])
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        preview = client.post(
            "/api/import/preview?mode=replace&target=device",
            files={"file": ("points.xlsx", stream.getvalue())},
        )
        assert preview.status_code == 200, preview.text
        data = preview.json()
        assert not data["errors"] and data["deleted"] == 0
        assert data["added"] == data["changed"] == 1
        point = data["config"]["devices"][0]["points"][1]
        assert client.post("/api/import/apply", json=data["config"]).status_code == 200
        actual = client.get("/api/config").json()["devices"][0]["points"]
        assert actual[0]["id"] == "old" and actual[1]["initial"] == point["initial"]
        assert client.post("/api/import/apply", json=data["config"]).status_code == 409
